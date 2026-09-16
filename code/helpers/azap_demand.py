# code/helpers/azap_demand.py — "New Export AZAP MMDDYY.xlsx" -> demand_plan_summary.csv
#
# The planners' weekly demand workbook (fs_manual/"New Export AZAP 091126.xlsx",
# ~14 MB, replaced every week, sometimes saved as .xlsm) carries the AZAP
# production plan for BOTH factories on the sheet 'pdp export AZAP n20'. The
# app consumes the small summary demand_plan_summary.csv (Week, Product, Tons)
# that used to be cut from it by hand. This module rebuilds that csv the way
# the hand-made one was, pinned against the drop of 2026-09-15: every one of
# its 238 rows reproduces exactly (tests/test_azap_demand.py).
#
# The recipe (drop of 2026-09-15):
#   * sheet 'pdp export AZAP n20', header in row 1, columns located by name
#     (case-insensitive, stripped): Factory, Product, Tons, Start Date;
#   * keep Factory == "NPA" only — NO 'To Export' filter (a "No" row still
#     counts, the hand-made file never filtered on it);
#   * drop non-numeric products (only "Trial" exists), skip rows with a blank
#     Tons, sum Tons per (ISO year, ISO week, Product) across machines;
#   * window = `weeks` (7) ISO weeks starting the first Monday STRICTLY after
#     the export date, which is the MMDDYY right after the prefix in the file
#     name ("091126" = 2026-09-11; copies such as "... 091126 (1).xlsx" or a
#     OneDrive conflict copy "... 091126-DESKTOP-X.xlsx" keep their date,
#     2026-09-16; fallback: the file's mtime date, reported as
#     export_date_source "mtime" so callers can warn);
#   * rows in window order (chronological, so a New-Year wrap keeps 52 before
#     1) then Product (6-digit code, zero-padded); csv = utf-8 with BOM, CRLF,
#     header "Week,Product,Tons", Tons with up to 3 decimals and trailing zeros
#     stripped ("39.003", "50", "8.999").
#
# Freshness (refresh_demand_summary): the rebuilt csv takes the workbook's
# mtime, so the live-data sync's settle rule (a file written seconds ago
# waits a pass) does not hold it back and the health page ages the demand by
# the export's own write time. Rule of 2026-09-16 (revised): the csv is "up
# to date" only when that stamp says it was built from the SELECTED workbook
# (same mtime, see _same_stamp), or when it is newer than EVERY AZAP
# workbook in the folder (a hand edit made after all of them). Comparing it
# with the selected workbook alone missed this week's export whenever the csv
# carried a later stamp from last week's workbook (an AutoSave rebuilt it,
# then the new export was copied in with its older mtime). So a hand edit
# wins until any AZAP workbook in the folder is written after it.
#
# Discovery (find_azap_workbook / passed_over_workbooks, 2026-09-16): the
# export date ranks the workbooks, the mtime breaks ties. A name date after
# the file's own save date (+1 day) cannot be the export date: that file
# ranks by its save date. A workbook saved after the export in use that
# loses the ranking is never silent: a name date more than four weeks before
# its save date (a year typo, "091825" for "091826", or an old export saved
# again) is a problem for the sync, any other one a note.
# The sync (scripts/fs-live-pull.py, folder mode) calls refresh on every
# source folder that holds a workbook; scripts/fs-live-push.py (work PC,
# GitHub mode) does the same best-effort when it runs from a repo checkout;
# scripts/azap_demand_summary.py is the hand-run CLI for the same thing.
#
# Write guards (2026-09-16): the csv lives in the SHARED fs_manual folder and
# syncs to SharePoint, so a bad build must never replace the good file. A
# build with 0 rows (wrong factory, a year typo in the name -> a window with
# no data) is refused with a ValueError naming the workbook, the window and
# the factory filter; so is a build with fewer than half the rows of the csv
# it would replace, unless force=True. The previous csv survives, its stamp
# still does not match the workbook, and the sync reports the refusal on
# every pass.

from __future__ import annotations

import os
import re
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd

SHEET_NAME = "pdp export AZAP n20"
CSV_NAME = "demand_plan_summary.csv"
WORKBOOK_GLOB = "New Export AZAP*.xls[xm]"
DEFAULT_WEEKS = 7
DEFAULT_FACTORY = "NPA"
REQUIRED_COLUMNS = ("Factory", "Product", "Tons", "Start Date")
CSV_COLUMNS = ("Week", "Product", "Tons")
# a build below this share of the existing csv's rows is refused unless forced
MIN_ROWS_RATIO = 0.5
# the csv swap on Windows is refused while someone has the csv open (Excel):
# retried like fs-live-pull's _replace_with_retry
REPLACE_ATTEMPTS = 5
REPLACE_PAUSE_S = 0.2
# (2026-09-16) the build copies the workbook's st_mtime_ns onto the csv; a
# copy through OneDrive/SharePoint may round one side to whole seconds, so
# "built from this workbook" also accepts a whole-second stamp within this of
# the other (see _same_stamp) — never a plain window: a stale csv copied in
# the same batch as a new workbook by a tool that does not keep file times
# lands a fraction of a second from it and must still be rebuilt
STAMP_ROUNDING_NS = 1_000_000_000
# (2026-09-16) a name date this many days after the file's save date cannot
# be the export date (the file ranks by its save date instead) ...
NAME_DATE_AHEAD_DAYS = 1
# ... and one more than this many days before it is doubtful (a year typo
# like 091825 / 010526, or an old export saved again)
NAME_DATE_BEHIND_DAYS = 28

# "New Export AZAP 091126.xlsx" / ".xlsm" — the six digits are the export date
# as MMDDYY. Whitespace between the name and the digits is tolerated, and so
# is anything after the six digits (2026-09-16): "New Export AZAP 091126
# (1).xlsx" and OneDrive's "New Export AZAP 091126-DESKTOP-X.xlsx" used to
# fall back to the mtime date and shift the window silently. A seventh digit
# is not a date ("0911267" -> None).
_NAME_RE = re.compile(r"^New Export AZAP\s*(\d{6})(?!\d).*\.xls[xm]$", re.IGNORECASE)


# ── file discovery ───────────────────────────────────────────────────────────

def export_date_from_name(path: str | Path) -> date | None:
    """MMDDYY in "New Export AZAP MMDDYY.xls[xm]" -> date; None when the name
    carries no (valid) date."""
    m = _NAME_RE.match(Path(path).name)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%m%d%y").date()
    except ValueError:
        return None


def _workbook_candidates(folder: Path) -> list[tuple[Path, os.stat_result]]:
    """Every "New Export AZAP*.xlsx|xlsm" file in `folder` with its stat.
    Raises OSError when the folder itself cannot be listed."""
    out = []
    for p in folder.glob(WORKBOOK_GLOB):
        try:
            if not p.is_file():
                continue
            st = p.stat()
        except OSError:
            continue                # vanished / unreadable between glob and stat
        out.append((p, st))
    return out


def rank_date(path: str | Path, mtime: float) -> date:
    """The export date a workbook ranks by (2026-09-16): the name's MMDDYY,
    else the save (mtime) date — and the save date too when the name date is
    more than NAME_DATE_AHEAD_DAYS after it (a file cannot be exported after
    it was saved: a typo such as 091827 would outrank every later export)."""
    saved = date.fromtimestamp(mtime)
    named = export_date_from_name(path)
    if named is None or named > saved + timedelta(days=NAME_DATE_AHEAD_DAYS):
        return saved
    return named


def name_date_doubt(path: str | Path, mtime: float) -> str | None:
    """Why the MMDDYY in the name is unlikely to be the export date, None
    when it is plausible or the name has none (2026-09-16): after the save
    date (+NAME_DATE_AHEAD_DAYS), or more than NAME_DATE_BEHIND_DAYS before
    it (a year typo, or an old export saved again)."""
    named = export_date_from_name(path)
    if named is None:
        return None
    saved = date.fromtimestamp(mtime)
    if named > saved + timedelta(days=NAME_DATE_AHEAD_DAYS):
        return f"the date in its name ({named}) is after the file was saved ({saved})"
    if named < saved - timedelta(days=NAME_DATE_BEHIND_DAYS):
        return (f"the date in its name ({named}) is {(saved - named).days} days before "
                f"the file was saved ({saved})")
    return None


def _ranked(cands: list[tuple[Path, os.stat_result]]) -> list[tuple[Path, os.stat_result]]:
    """Oldest first, newest last: export date (rank_date), then mtime."""
    return sorted(cands, key=lambda ps: (rank_date(ps[0], ps[1].st_mtime), ps[1].st_mtime_ns))


def find_azap_workbook(folder: str | Path) -> Path | None:
    """Newest "New Export AZAP*.xlsx|xlsm" in `folder`, None when there is
    none or the folder is not reachable. Excel's "~$..." lock files never
    match the prefix, so an open workbook is not mistaken for one.

    Newest = the latest EXPORT date first (rank_date: the MMDDYY in the
    name, else the mtime date), then the latest mtime (2026-09-16): an
    AutoSave on last week's workbook bumps its mtime but must not beat this
    week's export. A workbook this passes over although it was saved later
    is reported by passed_over_workbooks, never dropped silently."""
    folder = Path(folder)
    try:
        if not folder.is_dir():
            return None
        ranked = _ranked(_workbook_candidates(folder))
        return ranked[-1][0] if ranked else None
    except OSError:
        return None


def passed_over_workbooks(folder: str | Path, chosen: str | Path | None = None, *,
                          settle_seconds: float = 0.0,
                          now: float | None = None) -> list[dict]:
    """AZAP workbooks in `folder` that lose the ranking to `chosen` (default:
    find_azap_workbook's pick) although they may be the newer export
    (2026-09-16). Each entry: {"workbook": Path, "problem": bool, "message"}.

    * problem — its name date is doubtful (name_date_doubt) and it was saved
      after the export date in use: a year typo ("091825" for "091826",
      "010526" for "010527") would otherwise leave last week's demand in
      place with nothing on screen;
    * note (problem False) — any other workbook saved after the one in use
      (an AutoSave on last week's export): expected, but said out loud.

    A workbook modified less than `settle_seconds` before `now` may still be
    being copied in: it waits a pass, like everything else the sync reads.
    Never raises; an unreachable folder has nothing to report."""
    folder = Path(folder)
    try:
        if not folder.is_dir():
            return []
        ranked = _ranked(_workbook_candidates(folder))
    except OSError:
        return []
    if not ranked:
        return []
    top = ranked[-1]
    if chosen is not None:
        match = [ps for ps in ranked if ps[0].name == Path(chosen).name]
        if not match:
            return []
        top = match[0]
    chosen_path, chosen_st = top
    chosen_rank = rank_date(chosen_path, chosen_st.st_mtime)
    chosen_saved = datetime.fromtimestamp(chosen_st.st_mtime)
    now = time.time() if now is None else now
    out: list[dict] = []
    for p, st in ranked:
        if p.name == chosen_path.name:
            continue
        if settle_seconds and abs(now - st.st_mtime) < settle_seconds:
            continue
        saved = datetime.fromtimestamp(st.st_mtime)
        doubt = name_date_doubt(p, st.st_mtime)
        if doubt and saved.date() > chosen_rank:
            out.append({"workbook": p, "problem": True, "message": (
                f"{p.name} was saved {saved:%Y-%m-%d %H:%M}, after the export in use "
                f"({chosen_path.name}, export date {chosen_rank}), but {doubt} - a typo in its "
                f"MMDDYY (the year?) or an old export saved again; it is NOT used: rename it "
                f"'New Export AZAP MMDDYY.xlsx' with its export date, or move it out of the folder")})
        elif not doubt and st.st_mtime_ns > chosen_st.st_mtime_ns:
            named = export_date_from_name(p)
            by = f"{named} from its name" if named else \
                f"{saved.date()} from its modified date (no MMDDYY in the name)"
            out.append({"workbook": p, "problem": False, "message": (
                f"{p.name} (saved {saved:%Y-%m-%d %H:%M}) is newer on disk than "
                f"{chosen_path.name} (saved {chosen_saved:%Y-%m-%d %H:%M}) but its export date "
                f"is older ({by}); the demand is built from {chosen_path.name}")})
    return out


def first_monday_after(d: date) -> date:
    """The first Monday STRICTLY after `d` (a Monday export -> the next one)."""
    return d + timedelta(days=7 - d.weekday())


# ── cell coercion ────────────────────────────────────────────────────────────

def _product_code(value: Any) -> str | None:
    """Product cell -> 6-digit code (zero-padded), None when not numeric.
    The workbook stores ints; a float that is whole ("280480.0") or a digit
    string is accepted, "Trial" (and any other label) is dropped."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        n = value
    elif isinstance(value, float):
        if value != value or not value.is_integer():   # NaN / fractional
            return None
        n = int(value)
    else:
        s = str(value).strip()
        if not s.isdigit():
            return None
        n = int(s)
    if n < 0:
        return None
    return f"{n:06d}"


def _tons(value: Any) -> float | None:
    """Tons cell -> float; None for blank / non-numeric (row skipped)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        t = float(value)
    else:
        s = str(value).strip().replace(",", "")
        if not s:
            return None
        try:
            t = float(s)
        except ValueError:
            return None
    return None if t != t else t


def _as_date(value: Any) -> date | None:
    """Start Date cell (datetime / date / string / Excel serial) -> date."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # an Excel serial that data_only did not resolve (rare)
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%Y %H:%M:%S",
                "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    ts = pd.to_datetime(s, errors="coerce")
    return None if pd.isna(ts) else ts.date()


def format_tons(t: float) -> str:
    """Up to 3 decimals, trailing zeros stripped: 39.003 / 50 / 8.999."""
    s = f"{float(t):.3f}".rstrip("0").rstrip(".")
    return "0" if s in ("", "-0") else s


# ── the build ────────────────────────────────────────────────────────────────

def _open_sheet(workbook_path: Path, sheet: str):
    from openpyxl import load_workbook  # lazy: only the build needs it

    wb = load_workbook(str(workbook_path), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"{workbook_path.name}: sheet '{sheet}' not found "
            f"(sheets: {', '.join(wb.sheetnames)})")
    return wb, wb[sheet]


def _header_index(header_row) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).strip().lower()
        if key and key not in out:          # the first of a repeated header wins
            out[key] = i
    return out


def build_demand_summary(
    workbook_path: str | Path,
    *,
    weeks: int = DEFAULT_WEEKS,
    first_monday: date | datetime | str | None = None,
    factory: str = DEFAULT_FACTORY,
    sheet: str = SHEET_NAME,
) -> tuple[pd.DataFrame, dict]:
    """Read the AZAP workbook and return (summary frame, meta).

    Frame columns: Week (int, ISO week), Product (6-digit str), Tons (float,
    3 decimals), one row per (week, product) in the window, ordered
    chronologically then by Product. `first_monday` overrides the window
    start (default: first Monday strictly after the export date read from
    the file name, else from the file's mtime). Raises ValueError when the
    sheet or a required column is missing.
    """
    workbook_path = Path(workbook_path)
    if weeks < 1:
        raise ValueError(f"weeks must be >= 1 (got {weeks})")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"AZAP workbook not found: {workbook_path}")

    # window
    export_date = export_date_from_name(workbook_path)
    export_date_source = "name"
    if first_monday is not None:
        fm = _as_date(first_monday)
        if fm is None:
            raise ValueError(f"first_monday not a date: {first_monday!r}")
        if fm.weekday() != 0:
            raise ValueError(f"first_monday {fm} is not a Monday")
        export_date_source = "given"
    else:
        if export_date is None:
            export_date = date.fromtimestamp(workbook_path.stat().st_mtime)
            export_date_source = "mtime"
        fm = first_monday_after(export_date)
    mondays = [fm + timedelta(weeks=i) for i in range(int(weeks))]
    window = {(m.isocalendar()[0], m.isocalendar()[1]): i for i, m in enumerate(mondays)}

    wb, ws = _open_sheet(workbook_path, sheet)
    try:
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError(f"{workbook_path.name}: sheet '{sheet}' is empty") from None
        idx = _header_index(header)
        missing = [c for c in REQUIRED_COLUMNS if c.lower() not in idx]
        if missing:
            raise ValueError(
                f"{workbook_path.name}: sheet '{sheet}' lacks column(s) "
                f"{', '.join(repr(c) for c in missing)} (header row 1: "
                f"{', '.join(str(h) for h in header if h is not None)})")
        i_fac, i_prod = idx["factory"], idx["product"]
        i_tons, i_start = idx["tons"], idx["start date"]
        want = str(factory).strip().lower()

        rows_in = rows_factory = rows_kept = 0
        blank_tons = bad_dates = 0
        dropped_products: dict[str, int] = {}
        agg: dict[tuple[int, str], float] = {}   # (window slot, product) -> tons
        for r in rows:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            rows_in += 1
            fac = r[i_fac] if i_fac < len(r) else None
            if fac is None or str(fac).strip().lower() != want:
                continue
            rows_factory += 1
            prod = _product_code(r[i_prod] if i_prod < len(r) else None)
            if prod is None:
                raw = r[i_prod] if i_prod < len(r) else None
                label = "(blank)" if raw is None or str(raw).strip() == "" else str(raw).strip()
                dropped_products[label] = dropped_products.get(label, 0) + 1
                continue
            t = _tons(r[i_tons] if i_tons < len(r) else None)
            if t is None:
                blank_tons += 1
                continue
            d = _as_date(r[i_start] if i_start < len(r) else None)
            if d is None:
                bad_dates += 1
                continue
            iso = d.isocalendar()
            slot = window.get((iso[0], iso[1]))
            if slot is None:
                continue
            rows_kept += 1
            agg[(slot, prod)] = agg.get((slot, prod), 0.0) + t
    finally:
        wb.close()

    out_rows = [
        {"Week": mondays[slot].isocalendar()[1], "Product": prod, "Tons": round(t, 3)}
        for (slot, prod), t in sorted(agg.items())
    ]
    df = pd.DataFrame(out_rows, columns=list(CSV_COLUMNS))
    df["Week"] = df["Week"].astype(int)
    df["Tons"] = df["Tons"].astype(float)
    meta = {
        "workbook": str(workbook_path),
        "sheet": sheet,
        "factory": factory,
        "export_date": export_date,
        "export_date_source": export_date_source,
        "first_monday": mondays[0],
        "last_monday": mondays[-1],
        "weeks": [m.isocalendar()[1] for m in mondays],
        "rows_in": rows_in,
        "rows_factory": rows_factory,
        "rows_kept": rows_kept,
        "rows_out": len(df),
        "blank_tons": blank_tons,
        "bad_dates": bad_dates,
        "dropped_products": dropped_products,
    }
    return df, meta


def summary_csv_text(df: pd.DataFrame) -> str:
    """The csv body exactly as the hand-made file had it (CRLF, no BOM here)."""
    lines = [",".join(CSV_COLUMNS)]
    for _, r in df.iterrows():
        lines.append(f"{int(r['Week'])},{str(r['Product'])},{format_tons(r['Tons'])}")
    return "\r\n".join(lines) + "\r\n"


def write_demand_summary(df: pd.DataFrame, csv_path: str | Path) -> Path:
    """Write the summary: utf-8 with BOM, CRLF, "Week,Product,Tons". Written
    to a temp name and swapped in so a reader never sees a half file.

    The swap is retried REPLACE_ATTEMPTS times on PermissionError (Windows
    refuses it while the csv is open in Excel); when it still fails, or the
    temp write itself fails, the temp file is removed before the error is
    re-raised (2026-09-16): a "demand_plan_summary.csv.tmp" left in the
    shared fs_manual folder syncs to SharePoint for everyone to see."""
    csv_path = Path(csv_path)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    data = b"\xef\xbb\xbf" + summary_csv_text(df).encode("utf-8")
    tmp = csv_path.with_name(csv_path.name + ".tmp")
    try:
        tmp.write_bytes(data)
        for i in range(REPLACE_ATTEMPTS):
            try:
                os.replace(tmp, csv_path)
                break
            except PermissionError:
                if i >= REPLACE_ATTEMPTS - 1:
                    raise
                time.sleep(REPLACE_PAUSE_S)
    except BaseException:
        try:
            tmp.unlink()
        except OSError:
            pass                    # never there, or already swapped in
        raise
    return csv_path


def existing_csv_rows(csv_path: str | Path) -> int | None:
    """Data rows (non-blank lines after the header) of the csv a build would
    replace; None when there is no such file or it cannot be read."""
    p = Path(csv_path)
    try:
        if not p.is_file():
            return None
        text = p.read_bytes().decode("utf-8-sig", errors="replace")
    except OSError:
        return None
    lines = [ln for ln in text.splitlines() if ln.strip()]
    return max(len(lines) - 1, 0)


def write_refusal(rows: int, meta: dict, csv_path: str | Path, *,
                  force: bool = False) -> str | None:
    """Why a build of `rows` rows must NOT replace `csv_path` (None = write).

    Rules of 2026-09-16: 0 rows is always refused (a wrong factory or a year
    typo in the name leaves a window with no data, and a header-only csv
    would silently wipe the demand); fewer than MIN_ROWS_RATIO of the
    existing csv's data rows is refused unless `force`."""
    csv_path = Path(csv_path)
    wb_name = Path(str(meta.get("workbook") or "?")).name
    weeks = meta.get("weeks") or []
    window = (f"W{weeks[0]}..W{weeks[-1]} ({meta.get('first_monday')} .. "
              f"{meta.get('last_monday')}, export date {meta.get('export_date')} from "
              f"{meta.get('export_date_source')})") if weeks else "(no window)"
    if rows <= 0:
        return (f"{wb_name}: 0 summary rows for window {window} with factory filter "
                f"Factory == {meta.get('factory')!r} ({meta.get('rows_in', 0)} sheet rows read, "
                f"{meta.get('rows_factory', 0)} of that factory, {meta.get('rows_kept', 0)} in the "
                f"window) - refusing to write {csv_path.name}; the previous csv is kept")
    prev = existing_csv_rows(csv_path)
    if not force and prev and rows < MIN_ROWS_RATIO * prev:
        return (f"{wb_name}: {rows} summary rows for window {window} is below "
                f"{MIN_ROWS_RATIO:.0%} of the {prev} rows in the existing {csv_path.name} - "
                f"refusing to replace it; the previous csv is kept (force it with --force / "
                f"force=True when the smaller plan is real)")
    return None


def _same_stamp(csv_ns: int, wb_ns: int) -> bool:
    """The csv carries the workbook's modified time (2026-09-16): equal to
    the ns, or one side is a whole second within STAMP_ROUNDING_NS of the
    other (a sync that keeps only seconds rounded it). Two sub-second stamps
    that differ are different files' times, however close."""
    if csv_ns == wb_ns:
        return True
    return any(x % STAMP_ROUNDING_NS == 0 and abs(x - y) < STAMP_ROUNDING_NS
               for x, y in ((csv_ns, wb_ns), (wb_ns, csv_ns)))


def refresh_demand_summary(
    folder: str | Path,
    *,
    weeks: int = DEFAULT_WEEKS,
    force: bool = False,
    out_path: str | Path | None = None,
) -> dict:
    """Rebuild <folder>/demand_plan_summary.csv from the newest AZAP workbook
    in `folder` unless the csv is up to date.

    Up to date (2026-09-16, see the module header): the csv's mtime matches
    the selected workbook's (_same_stamp: it was built from it), or
    the csv is newer than EVERY AZAP workbook in `folder` (a hand edit made
    after all of them). Anything else rebuilds — a csv stamped by last
    week's AutoSaved workbook no longer hides this week's older-stamped
    export.

    Returns {ran, reason, workbook, csv, rows, meta, from_workbook}: ran
    False with reason "no workbook" / "up to date", ran True with reason
    "built" (or "forced"); from_workbook True when the csv now holds the
    selected workbook's build (built now, or its stamp matches), False when
    it is a hand edit newer than every workbook (or there is no workbook).
    A broken workbook raises (ValueError / openpyxl errors) — callers that
    must not die wrap it. The written csv takes the workbook's mtime (see the
    module header for why).

    Write guards (2026-09-16, see write_refusal): a build with 0 rows, or
    with fewer than half the rows of the csv it would replace (unless
    `force`), raises ValueError and writes nothing — the previous csv stays,
    its stamp still not the workbook's, so every later pass tries again and
    reports the refusal again. meta["export_date_source"] == "mtime" means
    the name carried no MMDDYY: callers surface it as a warning (and
    workbook_date_warning repeats it on the passes that do not rebuild).
    """
    folder = Path(folder)
    out = Path(out_path) if out_path else folder / CSV_NAME
    wb = find_azap_workbook(folder)
    if wb is None:
        return {"ran": False, "reason": f"no AZAP workbook in {folder}",
                "workbook": None, "csv": str(out), "rows": 0, "meta": None,
                "from_workbook": False}
    wb_stat = wb.stat()
    if not force and out.is_file():
        csv_ns = out.stat().st_mtime_ns
        from_wb = _same_stamp(csv_ns, wb_stat.st_mtime_ns)
        if not from_wb:
            try:
                newest_ns = max([st.st_mtime_ns for _, st in _workbook_candidates(folder)]
                                + [wb_stat.st_mtime_ns])
            except OSError:
                newest_ns = wb_stat.st_mtime_ns
        # "newer than every workbook" by a whole second: a csv stamped by
        # another workbook and rounded up by a sync is not a hand edit
        if from_wb or csv_ns - newest_ns >= STAMP_ROUNDING_NS:
            return {"ran": False, "reason": "up to date", "workbook": str(wb),
                    "csv": str(out), "rows": 0, "meta": None, "from_workbook": from_wb}
    df, meta = build_demand_summary(wb, weeks=weeks)
    refusal = write_refusal(len(df), meta, out, force=force)
    if refusal:
        raise ValueError(refusal)
    write_demand_summary(df, out)
    try:
        os.utime(out, ns=(wb_stat.st_atime_ns, wb_stat.st_mtime_ns))
    except OSError:
        pass  # the csv is newer than the workbook anyway: the rule still holds
    return {"ran": True, "reason": "forced" if force else "built", "workbook": str(wb),
            "csv": str(out), "rows": int(len(df)), "meta": meta, "from_workbook": True}


def export_date_warning(meta: dict | None) -> str | None:
    """One-line warning when the window came from the file's mtime because
    the name carries no MMDDYY (2026-09-16): the window may be off by a week
    or more without anything else looking wrong. None otherwise."""
    if not meta or meta.get("export_date_source") != "mtime":
        return None
    weeks = meta.get("weeks") or []
    span = f"W{weeks[0]}..W{weeks[-1]}" if weeks else "the window"
    return (f"{Path(str(meta.get('workbook') or '?')).name} has no MMDDYY date after "
            f"'New Export AZAP' in its name: {span} was taken from the file's modified date "
            f"{meta.get('export_date')} - rename it 'New Export AZAP MMDDYY.xlsx' if that is wrong")


def workbook_date_warning(workbook: str | Path | None, *,
                          weeks: int = DEFAULT_WEEKS) -> str | None:
    """export_date_warning for a workbook that was NOT rebuilt this pass
    (2026-09-16): the same text, with the window the build takes from the
    file's mtime, so a sync can repeat the warning on every pass while the
    csv built from that workbook is in use. None when the name carries a
    MMDDYY or the file cannot be read."""
    if workbook is None:
        return None
    p = Path(workbook)
    if export_date_from_name(p) is not None:
        return None
    try:
        saved = date.fromtimestamp(p.stat().st_mtime)
    except OSError:
        return None
    fm = first_monday_after(saved)
    return export_date_warning({
        "workbook": str(p), "export_date_source": "mtime", "export_date": saved,
        "weeks": [(fm + timedelta(weeks=i)).isocalendar()[1] for i in range(max(int(weeks), 1))]})
