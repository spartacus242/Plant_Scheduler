# code/stockcheck/x3_po_export.py — the ERP's purchase-order LINE export
# (order_npa.csv): Sage X3, one row per PO line, pipe-delimited, Windows-1252,
# NO quoting or escaping, LF or CRLF. It replaces IT's hand-run "NPA Open POs"
# workbook (2026-09-14). The whole layout is kept — 219 columns, every value,
# plus the untouched raw strings — so nothing the ERP sends is lost;
# to_po_lines() projects it onto the PoLine contract (po_import, plan §2)
# for the supply timeline.
#
# Layout, verified on the 2026-09-11 export (983 lines):
#   * delimiter '|', 219 fields on EVERY line (header + data), no quote
#     character: a literal '"' (the inch mark in a label description) and
#     commas (city "SOLON, OH") are plain text. Never hand this file to
#     csv.excel / csv.Sniffer — the lone '"' swallows the rest of the file;
#   * Windows-1252: é/è/ê/ô/à, « », ’ in the FRENCH header (the data rows of
#     the sample were pure ASCII; they are decoded the same way);
#   * dates are DD/MM/YYYY (1686 values with day > 12, none with month > 12);
#   * every quantity / price / amount is a 13-digit zero-padded integer with
#     FOUR implied decimals: '0000180000000' = 18000.0000 kg,
#     '0000000011463' = 1.1463 USD/KG — pinned against the old workbook on
#     three POs (qty, price and amount all agree; the ERP rounds amounts to
#     cents before scaling);
#   * (order_no, line_no) is unique; line_seq (the SECOND "Numéro de la
#     ligne") is the delivery sub-line. 15 French header names repeat, so
#     the layout maps duplicates by ORDER of occurrence;
#   * line_status (IT, 2026-09-15): 20 = receivable (pending receipt),
#     60 = archived (closed), 70 = deleted. A fully archived order drops out
#     of the export; an order keeps showing while some of its lines are
#     deleted and the rest not yet archived. So: 70 lines are `cancelled`
#     (never inbound, whatever their remaining qty), 60 lines are closed
#     (`received`), and "still inbound" is status 20 with qty_remaining > 0.
#     An unknown code falls back to the remaining-qty rule;
#   * order unit KEA = thousand each (1.5 KEA = 1,500 EA; the price is per
#     KEA). to_po_lines scales such lines to EA so the BOM join counts them;
#     purchasing still owes a written confirmation of the convention;
#   * the export leaves the ERP at ~19:10 and lands in the SharePoint library
#     "VIF Extracts" within the next 15-min refresh, same file name every
#     day, no archive copy.
#
# Best-effort like the other importers: problems land in errors, rows are
# kept (a mis-aligned row is flagged, never dropped), nothing raises.

from __future__ import annotations

import re
import time
import unicodedata
from collections import defaultdict, deque
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

ENCODING = "cp1252"
DELIM = "|"
SCALE = 10_000            # four implied decimals on every fixed-point field
LAYOUT_VERSION = "2026-09-11"
_MAX_ERRORS = 500         # the Inbound tab lists them; a broken file must not flood it

# kinds: T text (stripped, '' when blank) · D date DD/MM/YYYY -> ISO or None ·
# F fixed-point /10000 -> float or None · I integer -> int or None
T, D, F, I = "T", "D", "F", "I"

# (english key, French header as exported, kind) — in export order. Keys are
# the API; the French names are matched after accent/punctuation stripping,
# duplicates by occurrence. Empty-on-the-sample columns keep their slot.
COLUMNS: tuple[tuple[str, str, str], ...] = (
    ("company", "Société", T),
    ("company_name", "Raison sociale", T),
    ("site", "Etablissement", T),
    ("site_name", "Raison sociale", T),
    ("order_type", "Référence commande (origine = '1ACDE')", T),
    ("order_no", "Numéro de commande", T),
    ("supplier_order_ref", "Référence de la commande fournisseur", T),
    ("supplier_code", "Code fournisseur", T),
    ("partner_type", "Nature du tiers", T),
    ("supplier_name", "Nom du fournisseur", T),
    ("supplier_postcode", "Code postal du fournisseur", T),
    ("supplier_city", "Ville du fournisseur", T),
    ("supplier_country", "Pays du fournisseur", T),
    ("supplier_country_name", "Nom du pays", T),
    ("supplier_family", "Famille élémentaire du fournisseur", T),
    ("supplier_family_name", "Libellé de la famille élémentaire", T),
    ("supplier_acct_family", "Famille comptable du fournisseur", T),
    ("supplier_acct_family_name", "Libellé de la famille comptable", T),
    ("intermediary_company", "Société de l’intermédiaire", T),
    ("intermediary_company_name", "Raison sociale de la société de l’intermédiaire", T),
    ("intermediary_code", "Code de l’intermédiaire", T),
    ("intermediary_name", "Nom de l’intermédiaire", T),
    ("intermediary_postcode", "Code postal de l’intermédiaire", T),
    ("intermediary_city", "Ville de l’intermédiaire", T),
    ("intermediary_country", "Pays de l’intermédiaire", T),
    ("intermediary_country_name", "Libellé du pays de l’intermédiaire", T),
    ("order_contact", "Contact de la commande", T),
    ("order_contact_name", "Nom du contact", T),
    ("order_date", "Date de la commande", D),
    ("buyer", "Acheteur", T),
    ("buyer_name", "Nom de l’acheteur", T),
    ("planner", "Approvisionneur", T),
    ("planner_name", "Nom de l’approvisionneur", T),
    ("delivery_mode", "Mode de livraison", T),
    ("delivery_mode_name", "Libellé du mode de livraison", T),
    ("delivery_terms", "Conditions de livraison", T),
    ("delivery_terms_name", "Libellé des conditions de livraison", T),
    ("currency_mgmt", "Monnaie de gestion", T),
    ("currency_order", "Monnaie de la commande", T),
    ("currency_stat", "Monnaie de statistique", T),
    ("line_no", "Numéro de la ligne", T),
    ("line_direction", "Sens de la ligne", I),
    ("item", "Article commandé", T),
    ("item_name", "Libellé de l’article", T),
    ("item_family", "Famille élémentaire de l’article", T),
    ("item_family_name", "Libellé de la famille élémentaire", T),
    ("item_acct_family", "Famille comptable de l’article", T),
    ("item_acct_family_name", "Libellé de la famille comptable", T),
    ("contract_company", "Société du marché/offre", T),
    ("contract_company_name", "Libellé de la société du marché/offre", T),
    ("contract_site", "Etablissement du marché/offre", T),
    ("contract_site_name", "Libellé de l’établissement du marché/offre", T),
    ("contract_type", "Référence du marché/offre", T),
    ("contract_no", "Numéro interne du marché/offre", T),
    ("contract_line_no", "Numéro interne de la ligne du marché/offre", T),
    ("gross_price", "Prix brut", F),
    ("gross_price_2", "Prix brut", F),
    ("price_unit", "Unité de prix", T),
    ("special_price_flag", "Témoin prix spécial", I),
    ("gross_amount_mgmt", "Montant brut en monnaie de gestion", F),
    ("gross_amount_inv", "Montant brut en monnaie de facturation", F),
    ("gross_amount_stat", "Montant brut en monnaie de statisitiques", F),
    ("invoice_discount_mgmt", "Montant des remises sur facture", F),
    ("invoice_discount_inv", "Montant des remises sur facture", F),
    ("invoice_discount_stat", "Montant des remises sur facture", F),
    ("invoice_promo_mgmt", "Montant des promotions sur facture", F),
    ("invoice_promo_inv", "Montant des promotions sur facture", F),
    ("invoice_promo_stat", "Montant des promotions sur facture", F),
    ("invoice_freight_mgmt", "Montant des frais de transport sur facture", F),
    ("invoice_freight_inv", "Montant des frais de transport sur facture", F),
    ("invoice_freight_stat", "Montant des frais de transport sur facture", F),
    ("invoice_other_mgmt", "Montant des autres frais sur facture", F),
    ("invoice_other_inv", "Montant des autres frais sur facture", F),
    ("invoice_other_stat", "Montant des autres frais sur facture", F),
    ("offinvoice_discount_mgmt", "Montant des remises hors facture", F),
    ("offinvoice_discount_inv", "Montant des remises hors facture", F),
    ("offinvoice_discount_stat", "Montant des remises hors facture", F),
    ("offinvoice_promo_mgmt", "Montant des promotions hors facture", F),
    ("offinvoice_promo_inv", "Montant des promotions hors facture", F),
    ("offinvoice_promo_stat", "Montant des promotions hors facture", F),
    ("offinvoice_freight_mgmt", "Montant des frais de transport hors facture", F),
    ("offinvoice_freight_inv", "Montant des frais de transport hors facture", F),
    ("offinvoice_freight_stat", "Montant des frais de transport hors facture", F),
    ("offinvoice_other_mgmt", "Montant des autres frais hors facture", F),
    ("offinvoice_other_inv", "Montant des autres frais hors facture", F),
    ("offinvoice_other_stat", "Montant des autres frais hors facture", F),
    ("line_seq", "Numéro de la ligne", T),
    ("line_status", "Etat de la ligne", I),
    ("qty_ordered", "Quantité Commandée", F),
    ("qty_remaining", "Quantité restant en commande", F),
    ("order_unit", "Unité de commande", T),
    ("qty_ordered_stat", "Quantité Commandée", F),
    ("qty_remaining_stat", "Quantité restant en commande", F),
    ("stat_unit", "Unité de statistique", T),
    ("receipt_date", "Date de réception prévue", D),
    ("receipt_time", "Heure de réception prévue", T),
    ("receipt_company", "Société de réception", T),
    ("receipt_company_name", "Raison sociale de la société de réception", T),
    ("receipt_site", "Etablissement de réception", T),
    ("receipt_site_name", "Raison sociale de l’établissement de réception", T),
    ("receipt_location", "Lieu de réception", T),
    ("receipt_location_name", "Raison sociale du lieu de réception", T),
    ("receipt_postcode", "Code postal du lieu de réception", T),
    ("receipt_city", "Ville de réception", T),
    ("receipt_country", "Pays de réception", T),
    ("receipt_country_name", "Libellé du pays de réception", T),
    ("stock_company", "Société gérant les stocks", T),
    ("stock_company_name", "Raison sociale de la société gérant les stocks", T),
    ("stock_site", "Etablissement gérant les stocks", T),
    ("stock_site_name", "Raison sociale de l’établissement gérant les stocks", T),
    ("warehouse", "Dépôt de réception", T),
    ("warehouse_name", "Libellé du dépôt de réception", T),
    ("bin_location", "Emplacement de réception", T),
    ("bin_location_name", "Libellé de l’emplacement", T),
    ("analytic_section", "Section analytique", T),
    ("analytic_section_name", "Libellé de la section analytique", T),
    ("planned_lot", "Lot prévisionnel", T),
    ("internal_supplier_flag", "Témoin fournisseur interne", I),
    ("payto_company", "Société du fournisseur à payer", T),
    ("payto_company_name", "Raison sociale de la société du fournisseur à payer", T),
    ("payto_code", "Code fournisseur à payer", T),
    ("payto_name", "Nom du fournisseur à payer", T),
    ("payto_postcode", "Code postal du fournisseur à payer", T),
    ("payto_city", "Ville du fournisseur à payer", T),
    ("payto_country", "Pays du fournisseur à payer", T),
    ("payto_country_name", "Libellé du pays du fournisseur à payer", T),
    ("payto_family", "Famille élémentaire du fournisseur à payer", T),
    ("payto_family_name", "Libellé de la famille élémentaire du fournisseur à payer", T),
    ("payto_acct_family", "Famille comptable du fournisseur à payer", T),
    ("payto_acct_family_name", "Libellé de la famille comptable du fournisseur à payer", T),
    ("shipfrom_company", "Société du « Livré par »", T),
    ("shipfrom_company_name", "Raison sociale de la société « Livré par »", T),
    ("shipfrom_code", "Code « Livré par »", T),
    ("shipfrom_name", "Nom du « Livré par »", T),
    ("shipfrom_postcode", "Code postal du « Livré par »", T),
    ("shipfrom_city", "Ville du « Livré par »", T),
    ("shipfrom_country", "Pays du « Livré par »", T),
    ("shipfrom_country_name", "Libellé du pays du « Livré par »", T),
    ("shipfrom_family", "Famille élémentaire du « Livré par »", T),
    ("shipfrom_family_name", "Libellé de la famille élémentaire du « Livré par »", T),
    ("shipfrom_acct_family", "Famille comptable du « Livré par »", T),
    ("shipfrom_acct_family_name", "Libellé de la famille comptable du « Livré par »", T),
    ("billfrom_company", "Société du « Facturé par »", T),
    ("billfrom_company_name", "Raison sociale de la société « Facturé par »", T),
    ("billfrom_code", "Code « Facturé par »", T),
    ("billfrom_name", "Nom du « Facturé par »", T),
    ("billfrom_postcode", "Code postal du « Facturé par »", T),
    ("billfrom_city", "Ville du « Facturé par »", T),
    ("billfrom_country", "Pays du « Facturé par »", T),
    ("billfrom_country_name", "Libellé du pays du « Facturé par »", T),
    ("billfrom_family", "Famille élémentaire du « Facturé par »", T),
    ("billfrom_family_name", "Libellé de la famille élémentaire du « Facturé par »", T),
    ("billfrom_acct_family", "Famille comptable du « Facturé par »", T),
    ("billfrom_acct_family_name", "Libellé de la famille comptable du « Facturé par »", T),
    ("item_short_name", "Libellé réduit de l’article", T),
    ("item_ext_code_1", "Code externe 1 de l’article (ancien code)", T),
    ("item_ext_code_2", "Code externe 2 de l’article", T),
    ("item_ext_code_3", "Code externe 3 de l’article", T),
    ("supplier_ext_code_1", "Code externe 1 du fournisseur (Ancien code)", T),
    ("billing_company", "Code société de facturation", T),
    ("billing_company_name", "Raison sociale de la société de facturation", T),
    ("billing_site", "Etablissement de facturation", T),
    ("billing_site_name", "Raison sociale de l’établissement de facturation", T),
    ("requested_date", "Date d’arrivée demandée initialement", D),
    ("payment_terms", "Code du règlement", T),
    ("payment_terms_name", "Libellé du règlement", T),
    ("invoice_discount_sign", "Signe des remises sur facture", I),
    ("invoice_promo_sign", "Signe des promotions sur facture", I),
    ("invoice_freight_sign", "Signe des frais de transport sur facture", I),
    ("invoice_other_sign", "Signe des autres frais sur facture", I),
    ("offinvoice_discount_sign", "Signe des remises hors facture", I),
    ("offinvoice_promo_sign", "Signe des promotions hors facture", I),
    ("offinvoice_freight_sign", "Signe des frais de transport hors facture", I),
    ("offinvoice_other_sign", "Signe des autres frais hors facture", I),
    ("line_comment_internal", "Commentaire interne de la ligne", T),
    ("line_comment_external", "Commentaire externe de la ligne", T),
    ("contract_amount_mgmt", "Montant brut issu de l'offre/contrat en monnaie de gestion", F),
    ("contract_amount_order", "Montant brut issu de l'offre/contrat en monnaie de la commande", F),
    ("contract_amount_stat", "Montant brut issu de l'offre/contrat en monnaie statistique", F),
) + tuple(
    (f"header_crit_{kind}_{n}", f"{fr} critère entête {n}", T)
    for n in range(1, 11) for kind, fr in (("code", "Code"), ("value", "Valeur"))
) + tuple(
    (f"line_crit_{kind}_{n}", f"{fr} critère ligne {n}", T)
    for n in range(1, 11) for kind, fr in (("code", "Code"), ("value", "Valeur"))
)

KEYS: tuple[str, ...] = tuple(c[0] for c in COLUMNS)
KIND: dict[str, str] = {c[0]: c[2] for c in COLUMNS}
FRENCH: dict[str, str] = {c[0]: c[1] for c in COLUMNS}
assert len(COLUMNS) == 219 and len(set(KEYS)) == 219, "layout table drifted"

_DATE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_FIXED = re.compile(r"^-?\d+$")
_PROBE = "numero de commande"    # normalized header token that names the layout


@dataclass
class X3PoExport:
    """One parsed export. `rows` are decoded dicts keyed by KEYS (plus
    `extra_<i>` for header columns the layout does not know); `raw` holds
    the untouched strings of every data line in file order (the header is
    line 1, so raw[k] is file line k + 2) — nothing is thrown away."""
    keys: list[str] = field(default_factory=list)      # column keys in FILE order
    header: list[str] = field(default_factory=list)    # French names as read
    rows: list[dict] = field(default_factory=list)
    raw: list[list[str]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    source_path: str = ""
    source_mtime: str = ""
    n_rows: int = 0
    header_ok: bool = False
    unknown_columns: list[str] = field(default_factory=list)   # in file, not in layout
    missing_columns: list[str] = field(default_factory=list)   # in layout, not in file
    layout_version: str = LAYOUT_VERSION


# ── normalizing / decoding ─────────────────────────────────────────────────

def norm_header(name) -> str:
    """Accent-, case- and punctuation-insensitive header token: 'Nom du
    <Livré par »' (an ERP typo) and 'Nom du « Livré par »' both become
    'nom du livre par', so a quote-mark drift never breaks the mapping."""
    s = unicodedata.normalize("NFKD", str(name or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"[^0-9a-zA-Z]+", " ", s)
    return re.sub(r"\s+", " ", s).strip().casefold()


_EXPECTED_NORM: dict[str, str] = {k: norm_header(fr) for k, fr in FRENCH.items()}


def decode_text(raw: bytes) -> str:
    """Windows-1252 first (the ERP's code page: it has ’ « » that latin-1
    lacks); never raises."""
    try:
        return raw.decode(ENCODING)
    except UnicodeDecodeError:
        return raw.decode(ENCODING, errors="replace")


def _date(s: str) -> str | None:
    m = _DATE.match(s.strip())
    if not m:
        return None
    d, mo, y = (int(g) for g in m.groups())
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


def _fixed(s: str) -> float | None:
    s = s.strip()
    if not _FIXED.match(s):
        return None
    return int(s) / SCALE


def _int(s: str) -> int | None:
    s = s.strip()
    return int(s) if _FIXED.match(s) else None


def decode_value(kind: str, s: str) -> tuple[object, bool]:
    """(value, parsed). '' is None for D/F/I and '' for T (like the xlsx
    importer's blank cells). A value that does not parse comes back as the
    stripped string with parsed=False — the caller reports it, the value is
    never silently dropped."""
    if kind == T:
        return s.strip(), True
    if not s.strip():
        return None, True
    v = _date(s) if kind == D else _fixed(s) if kind == F else _int(s)
    return (s.strip(), False) if v is None else (v, True)


def looks_like_x3_export(path) -> bool:
    """True when the file's first line is the pipe-delimited Sage X3 PO-line
    header (≥ 50 pipes and the 'Numéro de commande' column). Reads only the
    head of the file; any read problem is False (the legacy loader then
    reports it its own way)."""
    try:
        with open(path, "rb") as f:
            head = f.read(65536)
    except OSError:
        return False
    first = decode_text(head).split("\n", 1)[0]
    return first.count(DELIM) >= 50 and _PROBE in norm_header(first)


# ── header mapping ─────────────────────────────────────────────────────────

def map_header(header: list[str]) -> tuple[list[str], list[str], list[str], list[str]]:
    """(keys in file order, unknown header names, missing layout keys, notes).

    Names are matched after normalization, duplicates by order of occurrence
    (the k-th 'Prix brut' in the file is the k-th in the layout). A header
    that matches nothing but sits at the SAME position as a layout key that
    nothing else claimed is accepted by position (a renamed/typo-fixed
    column keeps working, with a note); anything else unknown is kept under
    'extra_<i>' so its values survive too."""
    pool: dict[str, deque] = defaultdict(deque)
    for k in KEYS:
        pool[_EXPECTED_NORM[k]].append(k)
    keys: list[str | None] = []
    for h in header:
        q = pool.get(norm_header(h))
        keys.append(q.popleft() if q else None)
    taken = set(k for k in keys if k)
    notes: list[str] = []
    unknown: list[str] = []
    for i, (h, k) in enumerate(zip(header, keys)):
        if k is not None:
            continue
        cand = KEYS[i] if i < len(KEYS) else None
        if cand is not None and cand not in taken:
            keys[i] = cand
            taken.add(cand)
            notes.append(f"column {i + 1} {h!r} taken as {cand!r} "
                         f"(expected {FRENCH[cand]!r}) by position")
        else:
            keys[i] = f"extra_{i}"
            unknown.append(h)
    missing = [k for k in KEYS if k not in taken]
    return [str(k) for k in keys], unknown, missing, notes


# ── reader ─────────────────────────────────────────────────────────────────

def read_x3_po_export(path) -> X3PoExport:
    """Parse the export. Never raises: a missing/unreadable/foreign file
    yields empty rows and an error; rows with the wrong field count are
    kept (padded or over-long, flagged) so the planner sees them."""
    res = X3PoExport(source_path="" if path is None else str(path))
    if path is None or not str(path):
        res.errors.append("no PO file configured")
        return res
    p = Path(path)
    if not p.is_file():
        res.errors.append(f"file not found: {p}")
        return res
    try:
        res.source_mtime = time.strftime(
            "%Y-%m-%d %H:%M:%S", time.localtime(p.stat().st_mtime))
    except OSError:
        pass
    try:
        text = decode_text(p.read_bytes())
    except OSError as e:
        res.errors.append(f"cannot read {p.name}: {e}")
        return res
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    while lines and lines[-1] == "":
        lines.pop()
    if not lines:
        res.errors.append("empty file")
        return res
    header = [h.strip() for h in lines[0].split(DELIM)]
    res.header = header
    if header.count("") == len(header) or _PROBE not in norm_header(lines[0]):
        res.errors.append("header row not found (need the pipe-delimited "
                          "Sage X3 PO-line export: 'Numéro de commande', "
                          "'Article commandé', 'Quantité restant en "
                          "commande', 'Date de réception prévue' …)")
        return res
    keys, unknown, missing, notes = map_header(header)
    res.keys = keys
    res.unknown_columns = unknown
    res.missing_columns = missing
    res.header_ok = not unknown and not missing and not notes
    res.errors.extend(notes)
    if unknown:
        res.errors.append(f"{len(unknown)} column(s) not in the {LAYOUT_VERSION} "
                          f"layout, kept as extra_<n>: " + ", ".join(unknown[:5])
                          + (" …" if len(unknown) > 5 else ""))
    if missing:
        res.errors.append(f"{len(missing)} layout column(s) absent from the file: "
                          + ", ".join(missing[:5]) + (" …" if len(missing) > 5 else ""))
    ncol = len(keys)
    n_err = 0

    def err(msg: str) -> None:
        nonlocal n_err
        n_err += 1
        if n_err <= _MAX_ERRORS:
            res.errors.append(msg)

    for r, line in enumerate(lines[1:], start=2):
        if line == "":
            continue                     # a blank line inside the file is noise
        cells = line.split(DELIM)
        res.raw.append(list(cells))          # untouched, whatever the count
        if len(cells) != ncol:
            err(f"row {r}: {len(cells)} fields, expected {ncol} — kept but "
                f"its columns may be shifted (a '|' inside a text field?)")
            cells = (cells + [""] * ncol)[:ncol] if len(cells) < ncol else cells
        row: dict = {}
        for k, s in zip(keys, cells):
            kind = KIND.get(k, T)
            v, ok = decode_value(kind, s)
            if not ok:
                err(f"row {r}: {k} {s.strip()!r} is not a "
                    f"{'date' if kind == D else 'number'}")
            row[k] = v
        if len(cells) > ncol:
            row["extra_overflow"] = cells[ncol:]
        for k in missing:
            row[k] = None if KIND[k] != T else ""
        res.rows.append(row)
    if n_err > _MAX_ERRORS:
        res.errors.append(f"… and {n_err - _MAX_ERRORS} more row problem(s)")
    res.n_rows = len(res.rows)
    return res


# ── PoLine projection (po_import contract §2) ───────────────────────────────

def _slip(rd: str | None, ird: str | None) -> int | None:
    if not rd or not ird:
        return None
    try:
        return (date.fromisoformat(rd) - date.fromisoformat(ird)).days
    except ValueError:
        return None


def _f(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _crit(row: dict, code: str) -> str:
    """Value of the line criterion whose code is `code` ('PRO' producer,
    'COO' country of origin) — the ERP fills the criteria slots in a
    supplier-dependent order, so look them up by code, not slot."""
    for n in range(1, 11):
        if str(row.get(f"line_crit_code_{n}") or "").strip().upper() == code:
            return str(row.get(f"line_crit_value_{n}") or "").strip()
    return ""


# ERP line-status codes, confirmed by IT on 2026-09-15 (meeting on this export).
STATUS_RECEIVABLE = 20
STATUS_ARCHIVED = 60
STATUS_DELETED = 70
STATUS_TEXT = {STATUS_RECEIVABLE: "receivable", STATUS_ARCHIVED: "archived",
               STATUS_DELETED: "deleted"}
KEA_TO_EA = 1000.0     # order unit KEA = thousand each (IT, 2026-09-15)


def to_po_lines(export: X3PoExport) -> list[dict]:
    """PoLine dicts for the supply timeline — the 15 contract keys (see
    po_import) plus the export's own facts the planner asked to keep.

    qty = qty_remaining (what is still inbound; the received part of a
    partially delivered line is already in stock). received = nothing left
    to receive (qty_remaining ≤ 0) OR the line is archived (status 60).
    cancelled = the line is deleted in the ERP (status 70): the timeline
    gives it its own fate and never counts it, whatever remains on it.
    status_text names the code (receivable / archived / deleted, blank when
    unknown). A KEA line (thousand each) is scaled to EA, `unit_original`
    keeping "KEA". supplier_id drops the ERP's zero padding ('000048' ->
    '48', the id the old workbook showed); `supplier_code` keeps the exact
    ERP key."""
    from .po_import import po8   # lazy: po_import imports this module
    out: list[dict] = []
    for i, row in enumerate(export.rows):
        rd = row.get("receipt_date") if isinstance(row.get("receipt_date"), str) \
            and _DATE_ISO.match(row["receipt_date"]) else None
        ird = row.get("requested_date") if isinstance(row.get("requested_date"), str) \
            and _DATE_ISO.match(row["requested_date"]) else None
        od = row.get("order_date") if isinstance(row.get("order_date"), str) \
            and _DATE_ISO.match(row["order_date"]) else None
        rem = _f(row.get("qty_remaining"))
        ordered = _f(row.get("qty_ordered"))
        order_no = str(row.get("order_no") or "").strip()
        code = str(row.get("supplier_code") or "").strip()
        parts = [str(row.get(k) or "").strip()
                 for k in ("company", "site", "order_type")] + [order_no]
        stat_unit = str(row.get("stat_unit") or "").strip().upper()
        unit = str(row.get("order_unit") or "").strip()
        unit_original = ""
        if unit.upper() == "KEA":
            unit_original, unit = unit, "EA"
            rem = rem * KEA_TO_EA if rem is not None else None
            ordered = ordered * KEA_TO_EA if ordered is not None else None
        status = row.get("line_status")
        status_code = status if isinstance(status, int) and not isinstance(status, bool) else None
        cancelled = status_code == STATUS_DELETED
        received = (rem is not None and rem <= 0) or status_code == STATUS_ARCHIVED
        rem_stat = _f(row.get("qty_remaining_stat"))
        # the ERP only converts KG/L lines into the statistical unit (KG);
        # EA/M2 lines carry a 0 there, which is "no figure", not "0 kg"
        if stat_unit != "KG" or (rem_stat == 0 and unit.upper() != "KG"):
            rem_stat = None
        out.append({
            # contract keys
            "po": " ".join(x for x in parts if x),
            "po8": po8(order_no),
            "item": str(row.get("item") or "").strip(),
            "designation": str(row.get("item_name") or "").strip(),
            "qty": rem if rem is not None else 0.0,
            "unit": unit,
            "receipt_date": rd,
            "initial_receipt_date": ird,
            "slip_days": _slip(rd, ird),
            "arrival_area": str(row.get("receipt_location") or "").strip(),
            "supplier": str(row.get("supplier_name") or "").strip(),
            "supplier_id": code.lstrip("0") or code,
            "received": received,
            "order_date": od,
            "row": i + 2,
            # export facts kept for the planner (additive; legacy lines lack them)
            "qty_ordered": ordered,
            "qty_remaining_kg": rem_stat if stat_unit == "KG" else None,
            "status": status,
            "status_text": STATUS_TEXT.get(status_code, ""),
            "cancelled": cancelled,
            "unit_original": unit_original,
            "line_no": str(row.get("line_no") or "").strip(),
            "line_seq": str(row.get("line_seq") or "").strip(),
            "supplier_code": code,
            "supplier_ref": str(row.get("supplier_order_ref") or "").strip(),
            "warehouse": str(row.get("warehouse") or "").strip(),
            "receipt_time": str(row.get("receipt_time") or "").strip(),
            "price": _f(row.get("gross_price")),
            "price_unit": str(row.get("price_unit") or "").strip(),
            "amount": _f(row.get("gross_amount_inv")),
            "currency": str(row.get("currency_order") or "").strip(),
            "buyer": str(row.get("buyer") or "").strip(),
            "planner": str(row.get("planner") or "").strip(),
            "delivery_terms": str(row.get("delivery_terms") or "").strip(),
            "contract": str(row.get("contract_no") or "").strip(),
            "producer": _crit(row, "PRO"),
            "origin": _crit(row, "COO"),
            "comment": str(row.get("line_comment_internal") or "").strip(),
            "comment_external": str(row.get("line_comment_external") or "").strip(),
        })
    return out


_DATE_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


# ── clean re-exports (Excel-openable, nothing lost) ─────────────────────────

def write_clean_csv(export: X3PoExport, path, *, french_names: bool = False) -> Path:
    """UTF-8 (BOM, so Excel reads the accents), comma-separated, quoted —
    every column, decoded values (ISO dates, real decimals). Header = the
    English keys, or the French names with `french_names`."""
    import csv
    p = Path(path)
    keys = export.keys
    header = [FRENCH.get(k, k) for k in keys] if french_names else keys
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        for row in export.rows:
            w.writerow(["" if row.get(k) is None else row.get(k) for k in keys])
    return p


def write_clean_xlsx(export: X3PoExport, path) -> Path:
    """Workbook with two sheets: 'PO lines' (English keys, typed cells:
    dates as dates, numbers as numbers, codes as text so '000048' keeps its
    zeros) and 'Columns' (key ↔ French header ↔ kind), so the file can be
    read in Excel without guessing the layout."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO lines"
    keys = export.keys
    ws.append(keys)
    for row in export.rows:
        cells = []
        for k in keys:
            v = row.get(k)
            kind = KIND.get(k, T)
            if kind == D and isinstance(v, str) and _DATE_ISO.match(v):
                v = datetime.fromisoformat(v)
            cells.append("" if v is None else v)
        ws.append(cells)
    for j, k in enumerate(keys, start=1):
        col = get_column_letter(j)
        if KIND.get(k) == D:
            for c in ws[col][1:]:
                c.number_format = "yyyy-mm-dd"
        ws.column_dimensions[col].width = min(40, max(10, len(k) + 2))
    ws.freeze_panes = "A2"
    meta = wb.create_sheet("Columns")
    meta.append(["#", "key", "French header (as exported)", "kind"])
    kinds = {T: "text", D: "date (DD/MM/YYYY in the export)",
             F: "fixed-point ÷10000", I: "integer"}
    for j, k in enumerate(keys, start=1):
        fr = export.header[j - 1] if j - 1 < len(export.header) else FRENCH.get(k, "")
        meta.append([j, k, fr, kinds.get(KIND.get(k, T), "text")])
    meta.column_dimensions["B"].width = 30
    meta.column_dimensions["C"].width = 60
    meta.column_dimensions["D"].width = 30
    p = Path(path)
    wb.save(p)
    return p
