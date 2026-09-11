# tests/test_stock_reports.py — component-level views over the stock report
# (helpers/stock_reports.py): the purchasing list, the board/demand tables
# and the Excel export must agree with the report dict they are built from.

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "code"))

from helpers import stock_reports as sr  # noqa: E402


def _item(item, need, avail, status, unit="Kg", alts=()):
    ratio = (avail / need) if need else float("inf")
    return {"item": item, "designation": f"desc {item}", "need": need,
            "unit": unit, "available_primary": avail, "available_total": avail,
            "ratio": ratio, "status": status, "note": "",
            "alternates": [{"item": a, "designation": "", "available": 0} for a in alts]}


def fake_report() -> dict:
    return {
        "schedule_view": [
            {"block_id": "b1", "sku": "280351", "line_id": 0, "line_name": "P09",
             "start_h": 10.0, "end_h": 20.0, "qty_kg": 5000, "cases": 100,
             "status": "AT_RISK", "unk": [], "cycles": [],
             "items": [_item("730009", 300.0, 252.0, "AT_RISK", alts=("730009-V",)),
                       _item("752744", 10.0, 50.0, "OK", unit="M2")]},
            {"block_id": "b2", "sku": "280480", "line_id": 1, "line_name": "P10",
             "start_h": 5.0, "end_h": 12.0, "qty_kg": 4000, "cases": 80,
             "status": "TIGHT", "unk": [], "cycles": [],
             "items": [_item("730009", 100.0, 252.0, "TIGHT"),
                       _item("755000", 20.0, 0.0, "NOT_TRACKED", unit="EA")]},
            {"block_id": "b3", "sku": "120430", "line_id": 2, "line_name": "P11",
             "start_h": 30.0, "end_h": 40.0, "qty_kg": 1000, "cases": 20,
             "status": "OK", "unk": [], "cycles": [],
             "items": [_item("752744", 5.0, 50.0, "OK", unit="M2")]},
        ],
        "demand_view": [
            {"order_id": "280351-W0", "sku": "280351", "week_index": 0,
             "target_kg": 20000, "cases": 400, "achievable_ratio": 0.6,
             "status": "DO_NOT_SCHEDULE", "unk": [],
             "constraining": [_item("730009", 420.0, 252.0, "AT_RISK")]},
            {"order_id": "120430-W1", "sku": "120430", "week_index": 1,
             "target_kg": 3000, "cases": 60, "achievable_ratio": 5.0,
             "status": "OK", "unk": [], "constraining": []},
        ],
        "item_reverse": {
            "730009": [{"sku": "280351", "week_index": 0, "need": 420.0, "unit": "Kg"},
                       {"sku": "280351", "week_index": 1, "need": 100.0, "unit": "Kg"}],
            "730009-V": [{"sku": "280351", "week_index": 0, "need": 420.0, "unit": "Kg",
                          "as_alternate_of": "730009"}],
            "752744": [{"sku": "120430", "week_index": 1, "need": 15.0, "unit": "M2"}],
        },
        "no_bom_skus": ["280464"],
        "unk": [{"sku": "280480", "activity": "X", "item": "Y"}],
        "source_files": {"ediact 3.csv": "2026-09-11 08:00:00"},
        "import_errors": [],
    }


def test_component_shortages_lists_short_items_worst_first():
    rows = sr.component_shortages(fake_report(), stamp=lambda h: f"h{h:.0f}")
    items = [r["Item"] for r in rows]
    # 730009 is short on the board (AT_RISK + TIGHT), 755000 is untracked;
    # 752744 is fine everywhere and must NOT appear.
    assert items[0] == "730009"
    assert "755000" in items
    assert "752744" not in items
    r = {x["Item"]: x for x in rows}["730009"]
    assert r["Board need"] == 400.0            # 300 + 100 over the two blocks
    assert r["Available"] == 252.0
    assert r["Blocks at risk"] == 2
    assert r["First run at risk"] == "P10 h5"  # earliest short block wins
    assert r["Demand need"] == 520.0           # primaries only, both weeks
    assert r["Alternates"] == "730009-V"
    assert "280351" in r["SKUs"] and "280480" in r["SKUs"]
    assert r["Status"] == "AT RISK"


def test_component_shortages_respects_week_filter():
    rows = sr.component_shortages(fake_report(), week_index=1)
    r = {x["Item"]: x for x in rows}["730009"]
    assert r["Demand need"] == 100.0


def test_block_and_demand_rows_sort_worst_first():
    rep = fake_report()
    b = sr.block_rows(rep, stamp=lambda h: f"h{h:.0f}")
    assert [r["SKU"] for r in b] == ["280351", "280480", "120430"]
    assert b[0]["Items flagged"] == 1 and "730009 84%" in b[0]["Constraints"]
    assert "755000 not in stock export" in b[1]["Constraints"]
    risky = sr.block_rows(rep, stamp=lambda h: f"h{h:.0f}", only_risk=True)
    assert len(risky) == 2
    d = sr.demand_rows(rep, None, week_label=lambda w: f"WW{w:02d}")
    assert d[0]["Status"] == "DO NOT SCHEDULE" and d[0]["Week"] == "WW00"
    assert d[0]["Coverage"] == "60%"
    d1 = sr.demand_rows(rep, 1, week_label=lambda w: f"WW{w:02d}")
    assert [r["SKU"] for r in d1] == ["120430"]


def test_empty_report_is_harmless():
    empty = {"schedule_view": [], "demand_view": [], "item_reverse": {},
             "no_bom_skus": [], "unk": [], "source_files": {}, "import_errors": []}
    assert sr.component_shortages(empty) == []
    assert sr.block_rows(empty, stamp=str) == []
    assert sr.demand_rows(empty, None, str) == []
    xl = sr.excel_report(summary=[("Blocks", 0)], board=[], demand=[], shortages=[],
                         appointments=[], no_bom=[], unk=[])
    assert xl[:2] == b"PK"


def test_excel_report_has_every_sheet():
    import openpyxl
    rep = fake_report()
    xl = sr.excel_report(
        summary=[("Blocks at risk", 2), ("Report computed", "2026-09-11 08:00")],
        board=sr.block_rows(rep, stamp=lambda h: f"h{h:.0f}"),
        demand=sr.demand_rows(rep, None, week_label=lambda w: f"WW{w:02d}"),
        shortages=sr.component_shortages(rep),
        appointments=[{"po": "30042199", "date": "2026-09-12", "time": "08:00", "category": "CAPS"}],
        no_bom=rep["no_bom_skus"], unk=rep["unk"])
    wb = openpyxl.load_workbook(BytesIO(xl))
    assert wb.sheetnames == ["Summary", "Board", "Demand plan", "Component shortages",
                             "Receiving", "Data quality"]
    ws = wb["Component shortages"]
    header = [c.value for c in ws[1]]
    assert "Item" in header and "Board coverage" in header and "_status" not in header
    assert ws.freeze_panes == "A2"
