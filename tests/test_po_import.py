# tests/test_po_import.py — open-PO extract parser on a SYNTHETIC workbook
# that mirrors the IT layout (five 'U' headers, int item codes, sentinel
# row). The live file is never read (supplier names stay out of git).

from __future__ import annotations

import csv
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "code"))

from stockcheck.po_import import (  # noqa: E402
    PoResult, decode_batch_date, load_open_pos, po8, resolve_item_key)

# Column order copied from the 08-24 extract; the order unit is the 'U'
# right after 'Qty ordered at the origin', the other four are price units.
HEADERS = [
    "Order date", "Receipt number", "Supplier", "Supplier designation",
    "Order number", "Received item designation", "Supplier ord/rec reference",
    "Initial Receipt Date", "Receipt date", "Ordered item",
    "Qty ordered at the origin", "U", "Received item", "Qty invoiced", "UC",
    "Actual price", "U", "Actual price 3", "U", "Actual amount", "U",
    "Actual amount 2", "U", "Arrival area", "Batch", "Contract/Offer",
    "Supplier item reference",
]
LINE_KEYS = {"po", "po8", "item", "designation", "qty", "unit", "receipt_date",
             "initial_receipt_date", "slip_days", "arrival_area", "supplier",
             "supplier_id", "received", "order_date", "row"}
BLANK20 = " " * 20   # ERP pads blank text cells with spaces


def _row(po, item, qty, unit, receipt, initial, *, designation="SLV 24x90",
         area="RP1", supplier_id=48, supplier="GRAPHIC PACKAGING",
         receipt_no=BLANK20, order_date=datetime(2026, 7, 24)):
    u = unit.strip()
    return [order_date, receipt_no, supplier_id, supplier + "   ", po,
            designation + "   ", "   ", initial, receipt, item, qty, unit,
            "   ", 0, unit, 1.5, "USD/" + u, 1.5, "USD/" + u, 100, "USD", 100,
            "USD", area, "   ", "M2 02 1AMAR 30002236", None]


ROWS = [
    # one PO, three lines, item 754751 repeated (two deliveries)
    _row("2 02 1ACDE 30043537", 754751, 67200, "EA ", datetime(2026, 9, 2),
         datetime(2026, 8, 27)),
    _row("2 02 1ACDE 30043537", 754752, 1000, "EA ", datetime(2026, 9, 3),
         datetime(2026, 9, 3), designation="SLV 24x125"),
    _row("2 02 1ACDE 30043537", 754751, 2000, "EA ", datetime(2026, 9, 5),
         datetime(2026, 9, 1)),
    # received line: Receipt number filled
    _row("2 02 1ACDE 30043400", 752420, 5000, "M2 ", datetime(2026, 8, 28),
         datetime(2026, 8, 28), receipt_no="R000123", designation="FILM"),
    # unjoinable tooling code
    _row("2 02 1ACDE 30043900", 899999, 1, "EA ", datetime(2026, 9, 10),
         datetime(2026, 9, 10), designation="GPI TOOLING"),
    # raw line with a suffixed BOM key and a float qty
    _row("2 02 1ACDE 30043901", 735009, 500.5, "Kg ", datetime(2026, 9, 4),
         datetime(2026, 8, 20), designation="PUREE", area="RB1",
         supplier_id=250, supplier="Citrus Products Inc."),
    # garbage: text qty and text date
    _row("2 02 1ACDE 30043902", "X1", "abc", "EA ", "soon", None,
         designation="GARBAGE"),
]
GARBAGE_ROW = 8   # sheet row of the garbage line (header is row 1)
SENTINEL = ["?"] + [None] * (len(HEADERS) - 1)


def _write_xlsx(path, rows, headers=HEADERS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NPA Open POs"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()
    return path


def _csv_cell(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else v


# ── load_open_pos on xlsx ──────────────────────────────────────────────────

def test_xlsx_lines_and_meta(tmp_path):
    p = _write_xlsx(tmp_path / "open_pos.xlsx", ROWS + [SENTINEL])
    res = load_open_pos(p)
    assert isinstance(res, PoResult)
    assert res.n_rows == len(res.lines) == 7   # sentinel dropped, garbage kept
    assert res.source_path == str(p)
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", res.source_mtime)
    assert res.as_of is None
    assert res.max_receipt_date == "2026-09-10"
    assert all(set(ln) == LINE_KEYS for ln in res.lines)
    assert res.lines[0] == {
        "po": "2 02 1ACDE 30043537", "po8": "30043537", "item": "754751",
        "designation": "SLV 24x90", "qty": 67200.0, "unit": "EA",
        "receipt_date": "2026-09-02", "initial_receipt_date": "2026-08-27",
        "slip_days": 6, "arrival_area": "RP1", "supplier": "GRAPHIC PACKAGING",
        "supplier_id": "48", "received": False, "order_date": "2026-07-24",
        "row": 2,
    }
    assert isinstance(res.lines[0]["qty"], float)
    assert all(isinstance(ln["item"], str) for ln in res.lines)


def test_xlsx_multi_line_po_keeps_repeated_item(tmp_path):
    res = load_open_pos(_write_xlsx(tmp_path / "a.xlsx", ROWS + [SENTINEL]))
    po_a = [ln for ln in res.lines if ln["po8"] == "30043537"]
    assert [ln["item"] for ln in po_a] == ["754751", "754752", "754751"]
    assert [ln["qty"] for ln in po_a] == [67200.0, 1000.0, 2000.0]
    assert [ln["row"] for ln in po_a] == [2, 3, 4]


def test_xlsx_received_unjoinable_and_raw_lines(tmp_path):
    res = load_open_pos(_write_xlsx(tmp_path / "a.xlsx", ROWS + [SENTINEL]))
    by_row = {ln["row"]: ln for ln in res.lines}
    assert by_row[5]["received"] is True and by_row[5]["item"] == "752420"
    assert by_row[5]["unit"] == "M2"
    assert by_row[6]["item"] == "899999" and by_row[6]["received"] is False
    raw = by_row[7]
    assert raw["item"] == "735009" and raw["unit"] == "Kg"
    assert raw["arrival_area"] == "RB1" and raw["qty"] == 500.5
    assert raw["supplier"] == "Citrus Products Inc." and raw["supplier_id"] == "250"
    assert raw["slip_days"] == 15


def test_xlsx_garbage_row_kept_with_errors(tmp_path):
    res = load_open_pos(_write_xlsx(tmp_path / "a.xlsx", ROWS + [SENTINEL]))
    g = next(ln for ln in res.lines if ln["row"] == GARBAGE_ROW)
    assert g["qty"] == 0.0 and g["receipt_date"] is None
    assert g["initial_receipt_date"] is None and g["slip_days"] is None
    assert g["item"] == "X1" and g["po8"] == "30043902"
    row_errs = [e for e in res.errors if e.startswith(f"row {GARBAGE_ROW}:")]
    assert any("qty" in e for e in row_errs)
    assert any("Receipt date" in e for e in row_errs)
    # only the garbage row complains
    assert all(e.startswith(f"row {GARBAGE_ROW}:") for e in res.errors)


def test_xlsx_unit_is_first_u_after_qty(tmp_path):
    # a price 'U' BEFORE qty and two after: the first one after qty wins
    headers = ["Order number", "Ordered item", "Receipt date", "U",
               "Qty ordered at the origin", "U", "U"]
    row = ["2 02 1ACDE 30043537", 754751, datetime(2026, 9, 2), "USD/EA",
           67200, "EA ", "USD"]
    res = load_open_pos(_write_xlsx(tmp_path / "u.xlsx", [row], headers))
    assert res.lines[0]["unit"] == "EA" and res.lines[0]["qty"] == 67200.0
    assert res.lines[0]["received"] is False and res.lines[0]["supplier"] == ""
    assert any("missing optional columns" in e for e in res.errors)
    # no 'U' after qty at all → unit '' plus one file-level error
    res2 = load_open_pos(_write_xlsx(tmp_path / "nou.xlsx", [row[:5]],
                                     headers[:5]))
    assert res2.lines[0]["unit"] == ""
    assert any("no unit column" in e for e in res2.errors)


def test_xlsx_filled_row_without_item_is_reported(tmp_path):
    bad = list(ROWS[0])
    bad[9] = None
    res = load_open_pos(_write_xlsx(tmp_path / "a.xlsx", [bad, SENTINEL]))
    assert res.lines == [] and res.n_rows == 0
    assert any("row 2" in e and "Ordered item" in e for e in res.errors)


def test_xlsx_as_of_title_row_and_column(tmp_path):
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["As of", datetime(2026, 8, 24, 6, 10)])
    ws.append(HEADERS)
    ws.append(ROWS[0])
    wb.save(p)
    wb.close()
    res = load_open_pos(p)
    assert res.as_of == "2026-08-24"
    assert res.n_rows == 1 and res.lines[0]["row"] == 3
    # column variant
    p2 = _write_xlsx(tmp_path / "c.xlsx", [ROWS[0] + [datetime(2026, 8, 25)]],
                     HEADERS + ["Report date"])
    assert load_open_pos(p2).as_of == "2026-08-25"


# ── csv / empty / bad files ────────────────────────────────────────────────

@pytest.mark.parametrize("enc", ["utf-8-sig", "cp1252"])
def test_csv_variant(tmp_path, enc):
    p = tmp_path / "open_pos.csv"
    rows = [[_csv_cell(v) for v in r] for r in ROWS]
    rows[5][5] = "PURÉE   "   # non-ASCII survives both encodings
    with open(p, "w", encoding=enc, newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(rows)
        w.writerow(["?"] + [""] * (len(HEADERS) - 1))
    res = load_open_pos(p)
    assert res.n_rows == 7 and res.max_receipt_date == "2026-09-10"
    first = res.lines[0]
    assert first["item"] == "754751" and first["po8"] == "30043537"
    assert first["qty"] == 67200.0 and first["unit"] == "EA"
    assert first["receipt_date"] == "2026-09-02" and first["slip_days"] == 6
    assert first["supplier_id"] == "48" and first["received"] is False
    assert res.lines[3]["received"] is True
    assert res.lines[5]["designation"] == "PURÉE"
    assert res.lines[6]["qty"] == 0.0 and res.lines[6]["receipt_date"] is None


def test_empty_workbook(tmp_path):
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)
    wb.close()
    res = load_open_pos(p)
    assert res.lines == [] and res.n_rows == 0
    assert res.max_receipt_date is None and res.as_of is None
    assert res.errors and "empty" in res.errors[0]


def test_missing_headers(tmp_path):
    p = _write_xlsx(tmp_path / "bad.xlsx", [[1, 2, 3]], ["Foo", "Bar", "Baz"])
    res = load_open_pos(p)
    assert res.lines == [] and res.errors
    p2 = _write_xlsx(tmp_path / "bad2.xlsx", [[1, 2]],
                     ["Order number", "Ordered item"])
    res2 = load_open_pos(p2)
    assert res2.lines == []
    assert any("missing columns" in e and "receipt date" in e for e in res2.errors)


def test_missing_file_and_none_never_raise(tmp_path):
    res = load_open_pos(tmp_path / "nope.xlsx")
    assert res.lines == [] and res.errors and res.source_mtime == ""
    assert load_open_pos(None).lines == []
    # a non-workbook with an xlsx suffix is a file problem, not an exception
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip")
    res3 = load_open_pos(junk)
    assert res3.lines == [] and any("cannot read" in e for e in res3.errors)


# ── helpers ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("raw, expected", [
    ("2 02 1ACDE 30043537", "30043537"),
    ("30043537", "30043537"),
    ("abc", ""),
    ("30043600\n30043662", "30043600"),   # multi-PO dock cell: first run
    (30043537, "30043537"),
    (None, ""),
    ("", ""),
])
def test_po8(raw, expected):
    assert po8(raw) == expected


@pytest.mark.parametrize("raw, expected", [
    ("N62420340", date(2026, 8, 30)),
    ("n62420340", date(2026, 8, 30)),
    (" N60010001 ", date(2026, 1, 1)),
    ("N53650002", date(2025, 12, 31)),
    ("N53660002", None),      # 2025 has no day 366
    ("N60000002", None),      # day 0
    ("S123456", None),
    ("123456", None),
    ("", None),
    (None, None),
])
def test_decode_batch_date(raw, expected):
    assert decode_batch_date(raw) == expected


def test_resolve_item_key():
    bom = {"754751", "735009-A", "730058-A", "730058-B", "BT001"}
    assert resolve_item_key("754751", bom) == "754751"
    assert resolve_item_key("735009", bom) == "735009-A"   # single suffixed
    assert resolve_item_key("730058", bom) is None         # ambiguous
    assert resolve_item_key("899999", bom) is None
    assert resolve_item_key("", bom) is None
    assert resolve_item_key(None, bom) is None
    assert resolve_item_key(735009, bom) == "735009-A"     # int code tolerated
    assert resolve_item_key("BT001", bom) == "BT001"
